"""Build the uncorrected significance table (Macro-F1, Positive-F1, UAR).

Reads the CSV written by ``tools/paired_significance.py --metric-table`` together
with its provenance sidecar, refuses any table that is not labelled
``correction: none``, and writes one reader-facing workbook:

* ``Significance (uncorrected)``: only the comparisons that reach uncorrected
  significance in at least one metric, with the three metrics side by side;
* ``All comparisons``: every comparison, with typed numbers for sorting.

The workbook never replaces the corrected report under ``--output-dir``.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

METRICS = ("macro_f1", "positive_f1", "macro_recall")
METRIC_LABELS = {"macro_f1": "Macro-F1", "positive_f1": "Positive-F1", "macro_recall": "UAR"}


class TableError(RuntimeError):
    """Raised when the source table or its provenance cannot be trusted."""


def read_metric_table(csv_path: Path, sidecar_path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if not csv_path.is_file():
        raise TableError(f"metric table not found: {csv_path}")
    if not sidecar_path.is_file():
        raise TableError(f"provenance sidecar not found: {sidecar_path}")
    metadata = json.loads(sidecar_path.read_text(encoding="utf-8"))
    if metadata.get("schema_version") != "audiollm.metric_table.v1":
        raise TableError("unsupported metric table schema version")
    if metadata.get("correction") != "none":
        raise TableError("this builder only accepts an uncorrected table (correction: none)")
    if not metadata.get("source_report_sha256") or not metadata.get("family_sha256"):
        raise TableError("sidecar is missing the source report or family hash")

    with csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        raise TableError(f"no rows in {csv_path}")

    seen = {row["metric"] for row in rows}
    missing = [metric for metric in METRICS if metric not in seen]
    if missing:
        raise TableError(f"table is missing metrics: {missing}")

    for row in rows:
        for field in ("observed_delta", "bootstrap_ci_low", "bootstrap_ci_high", "p_value"):
            if row[field] in ("", None):
                raise TableError(f"{row['comparison_id']} / {row['metric']}: missing {field}")
            row[field] = float(row[field])
        row["subjects"] = int(row["subjects"])
        row["uncorrected_significant"] = row["uncorrected_significant"] == "True"
    return rows, metadata


def by_comparison(rows: list[dict[str, Any]]) -> dict[str, dict[str, dict[str, Any]]]:
    """Group rows into comparison -> metric -> row, checking the shared identity fields."""
    grouped: dict[str, dict[str, dict[str, Any]]] = {}
    identity: dict[str, tuple] = {}
    for row in rows:
        key = row["comparison_id"]
        grouped.setdefault(key, {})[row["metric"]] = row
        current = (row["block"], row["dataset"], row["subjects"], row["seeds"])
        if key in identity and identity[key] != current:
            raise TableError(f"{key}: identity fields differ between metrics")
        identity[key] = current
    for key, metrics in grouped.items():
        if set(metrics) != set(METRICS):
            raise TableError(f"{key}: expected {list(METRICS)}, found {sorted(metrics)}")
    return grouped


def _format_p(p_value: float) -> str:
    return "p<0.0001" if p_value < 0.0001 else f"p={p_value:.4f}"


def _format_cell(row: dict[str, Any]) -> str:
    mark = "*" if row["uncorrected_significant"] else ""
    return (f"{row['observed_delta']:+.4f} "
            f"[{row['bootstrap_ci_low']:+.4f}, {row['bootstrap_ci_high']:+.4f}] "
            f"{_format_p(row['p_value'])}{mark}")


def build_workbook(rows: list[dict[str, Any]], metadata: dict[str, Any], out_path: Path) -> Path:
    grouped = by_comparison(rows)
    tests = len(rows)
    alpha = float(metadata["alpha"])
    expected = metadata.get("expected_false_positives_at_alpha")
    significant = sorted(
        (key for key, metrics in grouped.items() if any(m["uncorrected_significant"] for m in metrics.values())),
        key=lambda key: min(metrics["p_value"] for metrics in grouped[key].values()),
    )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Significance (uncorrected)"
    notes = [
        f"Uncorrected significance, alpha = {alpha}. No family-wise correction is applied.",
        f"{tests} tests over {len(grouped)} comparisons; "
        f"{sum(1 for row in rows if row['uncorrected_significant'])} rows reach p<{alpha}; "
        f"about {expected} would be expected from chance alone.",
        "delta = comparison minus baseline, pooled out-of-fold subject-level; "
        "a positive delta means the second-named configuration is ahead.",
        "The deck and workbook headline cells are unweighted fold means, which is a different aggregation.",
        "Only comparisons significant in at least one metric are listed; every comparison is on the second sheet.",
        "* marks the metric that reaches p<alpha.",
    ]
    for index, note in enumerate(notes, start=1):
        cell = summary.cell(row=index, column=1, value=note)
        cell.font = Font(italic=True, bold=index == 1)
    header_row = len(notes) + 2
    headers = ["Block", "Dataset", "Comparison", "n", *(METRIC_LABELS[m] for m in METRICS)]
    for column, header in enumerate(headers, start=1):
        cell = summary.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True)
    for offset, key in enumerate(significant, start=header_row + 1):
        metrics = grouped[key]
        first = metrics[METRICS[0]]
        summary.cell(row=offset, column=1, value=first["block"])
        summary.cell(row=offset, column=2, value=first["dataset"])
        summary.cell(row=offset, column=3, value=key)
        summary.cell(row=offset, column=4, value=first["subjects"])
        for index, metric in enumerate(METRICS):
            summary.cell(row=offset, column=5 + index, value=_format_cell(metrics[metric]))
    for column, width in enumerate([34, 18, 52, 6, 34, 34, 34], start=1):
        summary.column_dimensions[get_column_letter(column)].width = width

    full = workbook.create_sheet("All comparisons")
    full_headers = ["Block", "Dataset", "Comparison", "n", "Seeds", "Permutation", "Mcnemar p"]
    for metric in METRICS:
        label = METRIC_LABELS[metric]
        full_headers += [f"{label} delta", f"{label} CI low", f"{label} CI high", f"{label} p"]
    for column, header in enumerate(full_headers, start=1):
        full.cell(row=1, column=column, value=header).font = Font(bold=True)
    for offset, key in enumerate(sorted(grouped, key=lambda k: (grouped[k][METRICS[0]]["block"], k)), start=2):
        metrics = grouped[key]
        first = metrics[METRICS[0]]
        full.cell(row=offset, column=1, value=first["block"])
        full.cell(row=offset, column=2, value=first["dataset"])
        full.cell(row=offset, column=3, value=key)
        full.cell(row=offset, column=4, value=first["subjects"])
        full.cell(row=offset, column=5, value=first["seeds"])
        full.cell(row=offset, column=6, value=first["permutation_method"])
        mcnemar = first.get("mcnemar_p_value")
        full.cell(row=offset, column=7, value=float(mcnemar) if mcnemar not in (None, "") else None)
        for index, metric in enumerate(METRICS):
            row = metrics[metric]
            base = 8 + index * 4
            for step, (field, number_format) in enumerate((
                ("observed_delta", "+0.0000;-0.0000"),
                ("bootstrap_ci_low", "+0.0000;-0.0000"),
                ("bootstrap_ci_high", "+0.0000;-0.0000"),
                ("p_value", "0.0000"),
            )):
                cell = full.cell(row=offset, column=base + step, value=row[field])
                cell.number_format = number_format
        full.cell(row=offset, column=7).number_format = "0.0000"
    for column, width in enumerate([34, 18, 52, 6, 7, 24, 10] + [11] * 12, start=1):
        full.column_dimensions[get_column_letter(column)].width = width
    full.freeze_panes = "D2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    print(f"wrote {out_path}")
    print(f"  Significance (uncorrected): {len(significant)} comparisons of {len(grouped)}")
    print(f"  All comparisons: {len(grouped)} rows x {len(METRICS)} metrics")
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--table", type=Path, default=PROJECT_ROOT / "outputs/significance/metric_table.csv")
    parser.add_argument("--sidecar", type=Path, default=None,
                        help="provenance JSON; defaults to the table path with a .json suffix")
    parser.add_argument("--out", type=Path,
                        default=PROJECT_ROOT / "outputs/significance/significance_uncorrected.xlsx")
    args = parser.parse_args()

    sidecar = args.sidecar if args.sidecar is not None else args.table.with_suffix(".json")
    rows, metadata = read_metric_table(args.table, sidecar)
    build_workbook(rows, metadata, args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
