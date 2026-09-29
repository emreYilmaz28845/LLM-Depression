#!/usr/bin/env python3
"""Build the separate experiment workbook for the label-vocabulary campaign.

This workbook is campaign-owned. The canonical workbook
(``depression_results_clean.xlsx``) is never touched: no canonical label, cell or
provenance entry changes here.

Sheets:

* ``Fixed test`` - one row per model/modality cell and training seed, five arm
  columns, each cell holding ``Macro-F1 / Positive-F1 / UAR``;
* ``Seed summary`` - per cell and arm, the mean and ddof=1 standard deviation of
  each metric over the three training seeds;
* ``Contrasts`` - the prespecified paired contrasts with delta, the unadjusted
  95% subject-clustered bootstrap interval, the permutation p-value and the Holm
  p-value inside that metric's own 30-member family;
* ``McNemar`` - the supporting exact McNemar view, 90 rows, corrected separately;
* ``Provenance`` - run name, attempt id, job ids, deployment, config hash,
  prediction hash and lifecycle state for every analysed run;
* ``Jobs`` - the Slurm reconciliation rows when a reconciliation file is given;
* ``Notes`` - definitions, caveats and the analysis contract.

Every number comes from ``analysis.json``, which is produced only from locally
validated REPORTABLE attempts.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

DEFAULT_ANALYSIS = Path("outputs/qwen3_daic_label_vocab/analysis")
METRIC_LABELS = {"macro_f1": "Macro-F1", "positive_f1": "Positive-F1", "macro_recall": "UAR"}


def _header(worksheet, columns: list[str]) -> None:
    worksheet.append(columns)
    for index in range(1, len(columns) + 1):
        cell = worksheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"


def _autosize(worksheet) -> None:
    for index, column in enumerate(worksheet.iter_cols(min_row=1, max_row=worksheet.max_row), start=1):
        longest = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column)
        worksheet.column_dimensions[get_column_letter(index)].width = min(40, max(12, longest + 2))


def load_analysis(analysis_dir: Path) -> dict[str, Any]:
    path = analysis_dir / "analysis.json"
    if not path.is_file():
        raise SystemExit(f"missing analysis payload: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def sheet_fixed_test(workbook: Workbook, payload: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("Fixed test")
    arms = payload["family"]["arm_order"]
    metrics = payload["family"]["metrics"]
    _header(worksheet, ["model/modality cell", "training seed", *arms])
    index = {(row["cell"], row["arm"], int(row["seed"])): row for row in payload["seed_rows"]}
    cells = sorted({row["cell"] for row in payload["seed_rows"]})
    seeds = sorted({int(row["seed"]) for row in payload["seed_rows"]})
    for cell in cells:
        for seed in seeds:
            values = []
            for arm in arms:
                row = index[(cell, arm, seed)]
                values.append(" / ".join(f"{row[metric]:.4f}" for metric in metrics))
            worksheet.append([cell, seed, *values])
    worksheet.append([])
    worksheet.append(["Cells hold " + " / ".join(METRIC_LABELS[metric] for metric in metrics)])
    worksheet.append([
        "DAIC fixed test partition, fold 0, likelihood backend, subject-level aggregation, "
        "harmonized_all_windows_full_coverage view. INVALID predictions count as wrong."
    ])
    _autosize(worksheet)


def sheet_seed_summary(workbook: Workbook, payload: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("Seed summary")
    metrics = payload["family"]["metrics"]
    columns = ["model/modality cell", "arm"]
    for metric in metrics:
        columns += [f"{METRIC_LABELS[metric]} mean", f"{METRIC_LABELS[metric]} sd (ddof=1)"]
    columns.append("seeds")
    _header(worksheet, columns)
    index = {(row["cell"], row["arm"], row["metric"]): row for row in payload["summary_long"]}
    seen: set[tuple[str, str]] = set()
    for row in payload["summary_long"]:
        key = (row["cell"], row["arm"])
        if key in seen:
            continue
        seen.add(key)
        record = [row["cell"], row["arm"]]
        for metric in metrics:
            match = index[(row["cell"], row["arm"], metric)]
            record += [round(match["mean"], 6), round(match["std_ddof1"], 6)]
        record.append(row["n_seeds"])
        worksheet.append(record)
    _autosize(worksheet)


def sheet_contrasts(workbook: Workbook, payload: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("Contrasts")
    _header(
        worksheet,
        [
            "model/modality cell", "earlier arm", "later arm", "metric", "delta",
            "unadjusted 95% CI low", "unadjusted 95% CI high", "p (permutation)",
            "p (Holm, family of 30)", "alpha", "subjects", "subject-seed keys",
        ],
    )
    for row in payload["contrasts"]:
        worksheet.append([
            row["cell"], row["arm_earlier"], row["arm_later"], METRIC_LABELS.get(row["metric"], row["metric"]),
            round(row["delta"], 6), round(row["ci_low"], 6), round(row["ci_high"], 6),
            row["p_value"], row["p_value_holm"], row["alpha"], row["subjects"], row["keys"],
        ])
    worksheet.append([])
    worksheet.append([
        "Delta is always (later arm) minus (earlier arm) in the order "
        + ", ".join(payload["family"]["arm_order"])
        + ". Bootstrap intervals are unadjusted and are never a corrected decision."
    ])
    _autosize(worksheet)


def sheet_mcnemar(workbook: Workbook, payload: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("McNemar")
    _header(
        worksheet,
        [
            "model/modality cell", "earlier arm", "later arm", "training seed",
            "baseline-only correct (b)", "comparison-only correct (c)",
            "accuracy delta", "p (exact)", "p (Holm, family of 90)",
        ],
    )
    for row in payload["mcnemar"]:
        worksheet.append([
            row["cell"], row["arm_earlier"], row["arm_later"], row["seed"], row["b"], row["c"],
            row["accuracy_delta"], row["p_value"], row["p_value_holm"],
        ])
    _autosize(worksheet)


def sheet_provenance(workbook: Workbook, payload: dict[str, Any], run_provenance: dict[str, Any] | None = None) -> None:
    worksheet = workbook.create_sheet("Provenance")
    if run_provenance:
        _header(
            worksheet,
            [
                "key", "run name", "attempt id", "lifecycle state", "subjects",
                "strict Macro-F1", "strict Positive-F1", "strict UAR",
                "label vocabulary", "model backend", "model revision", "model path",
                "git commit", "deployment id", "deployed source sha256",
                "manifest hash", "split metadata hash",
                "fold dir", "run_config path", "run_config sha256",
                "predictions path", "predictions sha256",
                "train job id", "best_eval job id",
            ],
        )
        for key, record in sorted(run_provenance.get("runs", {}).items()):
            strict = record.get("strict_headline") or {}
            job_ids = record.get("job_ids") or {}
            worksheet.append([
                key,
                record.get("run_name"),
                record.get("attempt_id"),
                record.get("state"),
                record.get("subjects"),
                strictly(strict, "binary_strict_macro_f1"),
                strictly(strict, "binary_strict_positive_f1"),
                strictly(strict, "binary_strict_uar"),
                record.get("label_vocab_version"),
                record.get("model_backend"),
                record.get("model_revision"),
                record.get("model_path"),
                record.get("git_commit"),
                record.get("deployment_id"),
                record.get("deployed_source_sha256"),
                record.get("manifest_hash"),
                record.get("split_metadata_hash"),
                record.get("fold_dir"),
                record.get("run_config_path"),
                record.get("run_config_sha256"),
                record.get("predictions_path"),
                record.get("predictions_sha256"),
                (job_ids.get("train") or [None])[0],
                (job_ids.get("best_eval") or [None])[0],
            ])
        worksheet.append([])
        worksheet.append([
            "Run evidence comes from the frozen run_provenance payload written by "
            "tools/qwen3_daic_label_vocab_analysis.py --verify-records; recorded sha256 values are "
            "recomputed and compared against the run's own artifacts.json, and no run evidence is "
            "modified by the analysis or the workbook."
        ])
        _autosize(worksheet)
        return
    _header(
        worksheet,
        [
            "key", "run name", "attempt id", "lifecycle state", "subjects",
            "strict Macro-F1", "strict Positive-F1", "strict UAR", "predictions sha256",
        ],
    )
    for key, record in sorted(payload["provenance"].items()):
        strict = record.get("strict_headline") or {}
        worksheet.append([
            key, record.get("run_name"), record.get("attempt_id"), record.get("state"), record.get("subjects"),
            strictly(strict, "binary_strict_macro_f1"), strictly(strict, "binary_strict_positive_f1"),
            strictly(strict, "binary_strict_uar"), record.get("predictions_sha256"),
        ])
    _autosize(worksheet)


def strictly(strict: dict[str, Any], key: str) -> Any:
    value = strict.get(key)
    return round(value, 6) if isinstance(value, (int, float)) else value


def sheet_jobs(workbook: Workbook, jobs_path: Path | None) -> None:
    if jobs_path is None or not jobs_path.is_file():
        return
    rows = json.loads(jobs_path.read_text(encoding="utf-8"))
    worksheet = workbook.create_sheet("Jobs")
    _header(
        worksheet,
        [
            "attempt id", "run name", "job key", "Slurm job id", "scheduler state",
            "accounting state", "exit code", "elapsed", "classification",
        ],
    )
    for row in rows:
        worksheet.append([
            row.get("attempt_id"), row.get("run_name"), row.get("job_key"), row.get("slurm_job_id"),
            row.get("squeue_state") or row.get("state"), row.get("account_state"),
            row.get("exit_code"), row.get("elapsed"), row.get("classification"),
        ])
    _autosize(worksheet)


def sheet_notes(workbook: Workbook, payload: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("Notes")
    _header(worksheet, ["note"])
    family = payload["family"]
    iterations = family["iterations"]
    notes = [
        "Campaign qwen3_daic_label_vocab_v1 (Qwen3 DAIC label-vocabulary). This workbook is campaign-owned; "
        "the canonical workbook depression_results_clean.xlsx is unchanged.",
        f"Contrast family: {family['arm_order']} arms, all ten unordered pairs in each of three cells, delta "
        "always later minus earlier; frozen before the first production submission.",
        f"Correction: Holm inside each metric's own 30-member family, alpha {family['alpha']}; exact McNemar "
        "corrected separately as its own 90-test family.",
        f"Tests: two-sided subject-clustered paired prediction-swap permutation with "
        f"{iterations['permutation']:,} iterations and a label-stratified subject-clustered bootstrap with "
        f"{iterations['bootstrap']:,} iterations, analysis seed {iterations['analysis_seed']}.",
        f"Metrics are computed per training seed and then averaged over seeds ({family['seeds']}); seed rows are "
        "never concatenated into one population and no best seed is selected.",
        "The EN arm keeps the canonical answer-only English instruction and has no legend line, while the four "
        "short arms render the legend. Every contrast that involves EN therefore bundles a vocabulary change "
        "with a legend and instruction-style change.",
        "Only locally validated REPORTABLE attempts feed this workbook; smoke attempts and any superseded or "
        "failed attempt are excluded, and a value that cannot be locally verified is left blank with a reason.",
        "DAIC keeps its fixed split: fold 0 is the single official test partition and the seed mean is not a "
        "held-out estimate over folds.",
        f"Analysis iterations actually used: permutation {payload['family']['iterations']['permutation']:,}, "
        f"bootstrap {payload['family']['iterations']['bootstrap']:,}; overridden for development: "
        f"{payload['family']['iterations']['iterations_overridden']}.",
    ]
    for note in notes:
        worksheet.append([note])
    _autosize(worksheet)


def build(*, analysis_dir: Path, output: Path, jobs_path: Path | None, run_provenance_path: Path | None = None) -> Path:
    payload = load_analysis(analysis_dir)
    run_provenance = None
    if run_provenance_path is not None:
        if not run_provenance_path.is_file():
            raise SystemExit(f"missing run provenance payload: {run_provenance_path}")
        run_provenance = json.loads(run_provenance_path.read_text(encoding="utf-8"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet_fixed_test(workbook, payload)
    sheet_seed_summary(workbook, payload)
    sheet_contrasts(workbook, payload)
    sheet_mcnemar(workbook, payload)
    sheet_provenance(workbook, payload, run_provenance)
    sheet_jobs(workbook, jobs_path)
    sheet_notes(workbook, payload)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--analysis-dir", default=str(DEFAULT_ANALYSIS))
    parser.add_argument("--output", required=True)
    parser.add_argument("--jobs", default=None, help="JSON list of reconciled job rows")
    parser.add_argument(
        "--run-provenance",
        default=None,
        help="run_provenance.json from qwen3_daic_label_vocab_analysis.py --verify-records",
    )
    args = parser.parse_args(argv)
    path = build(
        analysis_dir=Path(args.analysis_dir),
        output=Path(args.output),
        jobs_path=Path(args.jobs) if args.jobs else None,
        run_provenance_path=Path(args.run_provenance) if args.run_provenance else None,
    )
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
