from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("build_clean_workbook", ROOT / "scripts/build_clean_workbook.py")
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MODULE)


def _report() -> dict:
    metric = {
        "permutation": {
            "observed_delta": 0.1,
            "method": "exact_subject_paired",
            "p_value": 0.01,
            "p_value_holm_joint_block": 0.03,
            "p_value_holm_metric_block": 0.01,
            "p_value_holm_global": 0.06,
        },
        "bootstrap": {"ci_low": 0.02, "ci_high": 0.18},
    }
    return {
        "schema_version": "audiollm.significance_report.v1",
        "analysis_status": "retrospective_exploratory",
        "alpha": 0.05,
        "metrics": ["macro_f1", "positive_f1", "macro_recall"],
        "results": {"blocks": [{
            "id": "route",
            "comparisons": [{
                "id": "D|TF vs LR", "dataset": "d", "n_subjects": 10, "n_seeds": 1,
                "metrics": {name: metric for name in ("macro_f1", "positive_f1", "macro_recall")},
                "mcnemar": {
                    "status": "tested", "baseline_only_correct": 1, "comparison_only_correct": 5,
                    "p_value": 0.03, "p_value_holm_block": 0.03,
                },
            }],
        }]},
    }


def test_build_significance_sheets(tmp_path: Path):
    path = tmp_path / "report.json"
    path.write_text(json.dumps(_report()), encoding="utf-8")
    wb = Workbook()
    wb.remove(wb.active)
    MODULE.build_significance_sheets(wb, path)
    assert wb.sheetnames == ["Significance Summary", "Significance Full"]
    assert wb["Significance Summary"]["B5"].value == "D|TF vs LR"
    assert wb["Significance Summary"]["E5"].value == pytest.approx(0.1)
    assert wb["Significance Full"].max_row >= 8


def test_significance_report_rejects_wrong_status(tmp_path: Path):
    payload = _report()
    payload["analysis_status"] = "confirmatory"
    path = tmp_path / "report.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="retrospective_exploratory"):
        MODULE._load_significance_report(path)
