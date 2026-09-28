"""Tests for the campaign workbook builder.

The payload is the same shape ``tools/qwen3_daic_label_vocab_analysis.py`` writes,
so the workbook sheets are verified without any real run.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from tools.build_qwen3_daic_label_vocab_workbook import build

CELLS = ["text_only_qwen38_27b", "audio_only_qwen3omni_30b_a3b", "audio_text_qwen3omni_30b_a3b"]
ARMS = ["ab", "01", "truefalse", "yesno", "en"]
METRICS = ["macro_f1", "positive_f1", "macro_recall"]
SEEDS = [7, 1337, 2024]


def _payload() -> dict:
    seed_rows = []
    summary = []
    for cell in CELLS:
        for arm in ARMS:
            values = {metric: [] for metric in METRICS}
            for seed in SEEDS:
                row = {"cell": cell, "model": "m", "modality": "x", "arm": arm, "seed": seed}
                for index, metric in enumerate(METRICS):
                    value = 0.5 + index * 0.01 + seed / 100000.0
                    row[metric] = value
                    values[metric].append(value)
                seed_rows.append(row)
            for metric in METRICS:
                mean = sum(values[metric]) / len(values[metric])
                summary.append(
                    {
                        "cell": cell, "model": "m", "modality": "x", "arm": arm, "metric": metric,
                        "mean": mean, "std_ddof1": 0.001, "n_seeds": 3,
                    }
                )
    contrasts = []
    mcnemar = []
    for index, cell in enumerate(CELLS):
        pairs = [(ARMS[a], ARMS[b]) for a in range(len(ARMS)) for b in range(a + 1, len(ARMS))]
        for earlier, later in pairs:
            for metric in METRICS:
                contrasts.append(
                    {
                        "cell": cell, "model": "m", "modality": "x", "arm_earlier": earlier, "arm_later": later,
                        "metric": metric, "delta": 0.01, "p_value": 0.2, "ci_low": -0.01, "ci_high": 0.03,
                        "subjects": 47, "keys": 141, "p_value_holm": 1.0, "alpha": 0.05,
                    }
                )
            for seed in SEEDS:
                mcnemar.append(
                    {
                        "cell": cell, "arm_earlier": earlier, "arm_later": later, "seed": seed,
                        "b": 3, "c": 5, "accuracy_delta": 0.02, "p_value": 0.7, "p_value_holm": 1.0,
                    }
                )
    provenance = {
        f"{cell}|{arm}|{seed}": {
            "run_name": f"run-{cell}-{arm}-{seed}", "attempt_id": f"att-{cell}-{arm}-{seed}",
            "state": "REPORTABLE", "subjects": 47, "predictions_sha256": "a" * 64,
            "strict_headline": {"binary_strict_macro_f1": 0.5, "binary_strict_positive_f1": 0.4, "binary_strict_uar": 0.55},
        }
        for cell in CELLS for arm in ARMS for seed in SEEDS
    }
    return {
        "schema_version": "audiollm.qwen3_daic_label_vocab.analysis.v1",
        "family": {
            "arm_order": ARMS, "seeds": SEEDS, "alpha": 0.05, "metrics": METRICS,
            "iterations": {"permutation": 1_000_000, "bootstrap": 100_000, "iterations_overridden": False, "analysis_seed": 1337},
        },
        "seed_rows": seed_rows,
        "summary_long": summary,
        "contrasts": contrasts,
        "mcnemar": mcnemar,
        "provenance": provenance,
    }


def test_workbook_has_the_expected_sheets_and_rows(tmp_path: Path) -> None:
    analysis = tmp_path / "analysis"
    analysis.mkdir()
    (analysis / "analysis.json").write_text(json.dumps(_payload()), encoding="utf-8")
    jobs = tmp_path / "jobs.json"
    jobs.write_text(
        json.dumps([{"attempt_id": "att-1", "run_name": "run-1", "job_key": "train", "slurm_job_id": "1", "state": "COMPLETED", "exit_code": "0:0"}]),
        encoding="utf-8",
    )

    output = build(analysis_dir=analysis, output=tmp_path / "workbook.xlsx", jobs_path=jobs)
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Fixed test", "Seed summary", "Contrasts", "McNemar", "Provenance", "Jobs", "Notes"]

    fixed = workbook["Fixed test"]
    assert fixed.cell(row=1, column=3).value == "ab"
    assert fixed.cell(row=2, column=2).value == 7
    assert " / " in str(fixed.cell(row=2, column=3).value)

    summary = workbook["Seed summary"]
    assert summary.max_row - 1 == len(CELLS) * len(ARMS)  # one row per cell and arm

    contrasts = workbook["Contrasts"]
    assert contrasts.max_row - 1 >= 90  # 30 contrasts x 3 metrics plus the footnote row

    mcnemar = workbook["McNemar"]
    assert mcnemar.max_row - 1 == 90

    provenance = workbook["Provenance"]
    assert provenance.max_row - 1 == len(CELLS) * len(ARMS) * len(SEEDS)
    assert provenance.cell(row=2, column=4).value == "REPORTABLE"

    jobs_sheet = workbook["Jobs"]
    assert jobs_sheet.cell(row=2, column=4).value == "1"

    notes = workbook["Notes"]
    text = "\n".join(str(cell.value) for cell in notes["A"])
    assert "canonical workbook depression_results_clean.xlsx is unchanged" in text
    assert "legend" in text
