"""Tests for the uncorrected significance workbook builder."""

from __future__ import annotations

import csv
import importlib.util as _ilu
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
_spec = _ilu.spec_from_file_location(
    "build_uncorrected_significance_table", ROOT / "tools/build_uncorrected_significance_table.py"
)
builder = _ilu.module_from_spec(_spec)
_spec.loader.exec_module(builder)

FIELDS = [
    "block", "comparison_id", "correction_family", "dataset", "subjects", "seeds", "metric",
    "observed_delta", "bootstrap_ci_low", "bootstrap_ci_high", "permutation_method", "p_value",
    "uncorrected_significant", "mcnemar_p_value",
]


def _row(comparison_id: str, metric: str, p_value: float) -> dict:
    return {
        "block": "synthetic_block",
        "comparison_id": comparison_id,
        "correction_family": "F1|backbone|dataset=d3tec",
        "dataset": "d3tec",
        "subjects": 62,
        "seeds": 1,
        "metric": metric,
        "observed_delta": -0.13 if p_value < 0.05 else 0.01,
        "bootstrap_ci_low": -0.2268,
        "bootstrap_ci_high": -0.0463,
        "permutation_method": "exact_subject_paired",
        "p_value": p_value,
        "uncorrected_significant": p_value < 0.05,
        "mcnemar_p_value": 0.0117,
    }


def _write_table(tmp_path: Path, correction: str = "none") -> tuple[Path, Path]:
    rows = []
    for metric, p_value in zip(builder.METRICS, (0.0068, 0.0052, 0.0118)):
        rows.append(_row("synthetic|A vs B", metric, p_value))
    for metric in builder.METRICS:
        rows.append(_row("synthetic|C vs D", metric, 0.62))
    table_path = tmp_path / "metric_table.csv"
    with table_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    sidecar_path = table_path.with_suffix(".json")
    sidecar_path.write_text(json.dumps({
        "schema_version": "audiollm.metric_table.v1",
        "correction": correction,
        "alpha": 0.05,
        "expected_false_positives_at_alpha": 0.3,
        "source_report_sha256": "a" * 64,
        "family_sha256": "b" * 64,
    }), encoding="utf-8")
    return table_path, sidecar_path


def test_build_workbook_lists_only_significant_comparisons_in_summary(tmp_path: Path):
    table_path, sidecar_path = _write_table(tmp_path)
    rows, metadata = builder.read_metric_table(table_path, sidecar_path)
    out_path = builder.build_workbook(rows, metadata, tmp_path / "out.xlsx")

    workbook = load_workbook(out_path)
    assert workbook.sheetnames == ["Significance (uncorrected)", "All comparisons"]
    summary = workbook["Significance (uncorrected)"]
    summary_text = "\n".join(str(cell.value) for row in summary.iter_rows() for cell in row if cell.value)
    assert "Uncorrected significance, alpha = 0.05" in summary_text
    assert "unweighted fold means" in summary_text
    assert "synthetic|A vs B" in summary_text
    assert "synthetic|C vs D" not in summary_text
    assert "-0.1300 [-0.2268, -0.0463] p=0.0068*" in summary_text

    full = workbook["All comparisons"]
    comparison_ids = [full.cell(row=index, column=3).value for index in range(2, full.max_row + 1)]
    assert sorted(comparison_ids) == ["synthetic|A vs B", "synthetic|C vs D"]
    assert full.cell(row=2, column=8).value == pytest.approx(-0.13)
    assert full.cell(row=2, column=7).value == pytest.approx(0.0117)


def test_read_metric_table_refuses_a_corrected_sidecar(tmp_path: Path):
    table_path, sidecar_path = _write_table(tmp_path, correction="holm")
    with pytest.raises(builder.TableError, match="uncorrected table"):
        builder.read_metric_table(table_path, sidecar_path)


def test_by_comparison_requires_every_metric(tmp_path: Path):
    rows = [_row("synthetic|A vs B", metric, 0.5) for metric in builder.METRICS[:2]]
    with pytest.raises(builder.TableError, match="expected"):
        builder.by_comparison(rows)
